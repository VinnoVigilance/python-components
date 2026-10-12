"""Prepare a body-based VV silver-label dataset without splitting or training.

Label values: 1 = both models positive, 0 = both models negative (provisional),
-1 = model disagreement (exclude this article/label pair from supervised loss).
Only `text` is a classifier feature; model tags and IDs are metadata/targets.

Requires audit_dataset.py in the same directory.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import tempfile

from openpyxl import load_workbook

try:
    from audit_dataset import LABEL_NAMES, audit, parse_tags, text_value
except ModuleNotFoundError as error:
    if error.name != "audit_dataset":
        raise
    raise SystemExit("ERROR: put audit_dataset.py next to prepare_dataset.py.") from error


PREPROCESSING_VERSION = "vv-consensus-body-v1"


def clean_text(value: object) -> str:
    text = text_value(value)
    return "" if text.casefold() == "nan" else re.sub(r"\s+", " ", text).strip()


def read_records(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = [text_value(value) for value in next(rows)]
        records = []
        for excel_row, values in enumerate(rows, start=2):
            if not any(value is not None for value in values):
                continue
            row = dict(zip(header, values))
            title = clean_text(row.get("News title"))
            body = clean_text(row.get("News body"))
            body_hash = hashlib.sha256(body.casefold().encode("utf-8")).hexdigest() if body else None
            records.append({
                "news_id": text_value(row.get("News ID")),
                "excel_row": excel_row,
                "title": title,
                "body": body,
                "body_hash": body_hash,
                "claude": parse_tags(row.get("Claude tags")),
                "gpt": parse_tags(row.get("GPT tags")),
            })
        return records
    finally:
        workbook.close()


def prepare(records: list[dict]) -> tuple[list[dict], list[dict], dict]:
    groups = defaultdict(list)
    excluded = []

    def exclude(record: dict, reason: str, **details: object) -> None:
        excluded.append({"news_id": record["news_id"], "excel_row": record["excel_row"],
                         "reason": reason, **details})

    for record in records:
        if not record["body"]:
            exclude(record, "missing_body")
        else:
            groups[record["body_hash"]].append(record)

    prepared = []
    for fingerprint, group in groups.items():
        annotations = {(tuple(sorted(row["claude"])), tuple(sorted(row["gpt"]))) for row in group}
        if len(annotations) > 1:
            for row in group:
                exclude(row, "duplicate_annotation_conflict", body_hash=fingerprint)
            continue
        representative = group[0]
        for extra in group[1:]:
            exclude(extra, "consistent_duplicate", kept_news_id=representative["news_id"],
                    body_hash=fingerprint)
        claude, gpt = representative["claude"], representative["gpt"]
        label_values = [
            1 if label in claude and label in gpt else
            -1 if (label in claude) != (label in gpt) else 0
            for label in LABEL_NAMES
        ]
        title, body = representative["title"], representative["body"]
        prepared.append({
            "news_id": representative["news_id"],
            "excel_row": representative["excel_row"],
            "original_news_ids": [row["news_id"] for row in group],
            "body_hash": fingerprint,
            "title": title,
            "body": body,
            "text": f"{title}\n{body}" if title else body,
            "claude_labels": [label for label in LABEL_NAMES if label in claude],
            "gpt_labels": [label for label in LABEL_NAMES if label in gpt],
            "shared_labels": [label for label in LABEL_NAMES if label in claude and label in gpt],
            "label_values": label_values,
        })
    prepared.sort(key=lambda row: row["excel_row"])
    excluded.sort(key=lambda row: row["excel_row"])
    exclusions = Counter(row["reason"] for row in excluded)
    statistics = {
        "input_rows": len(records),
        "prepared_rows": len(prepared),
        "excluded_rows": len(excluded),
        "excluded_by_reason": dict(exclusions),
        "rows_with_no_consensus_positive": sum(not row["shared_labels"] for row in prepared),
        "rows_with_any_unknown_label": sum(-1 in row["label_values"] for row in prepared),
        "label_counts": {
            label: {
                "consensus_positive": sum(row["label_values"][index] == 1 for row in prepared),
                "provisional_negative": sum(row["label_values"][index] == 0 for row in prepared),
                "unknown": sum(row["label_values"][index] == -1 for row in prepared),
            } for index, label in enumerate(LABEL_NAMES)
        },
    }
    if len(prepared) + len(excluded) != len(records):
        raise ValueError("Row accounting failed.")
    if len({row["body_hash"] for row in prepared}) != len(prepared):
        raise ValueError("Duplicate text remains in prepared dataset.")
    if not prepared:
        raise ValueError("No eligible body-based records remain.")
    return prepared, excluded, statistics


def write_atomic(path: Path, content: str) -> str:
    encoded = content.encode("utf-8")
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="wb", dir=path.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(encoded)
        os.replace(temporary, path)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return hashlib.sha256(encoded).hexdigest()


def json_lines(records: list[dict]) -> str:
    return "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in records)


def main() -> int:
    base = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=base / "data" / "raw" / "News.xlsx")
    parser.add_argument("--output-dir", type=Path, default=base / "data" / "processed")
    args = parser.parse_args()
    if not args.input.is_file():
        parser.error(f"Input file not found: {args.input}")
    try:
        source_audit = audit(args.input)
        blockers = (source_audit["missing_id_rows"] or source_audit["duplicate_news_ids"]
                    or source_audit["unknown_labels"] or source_audit["intersection_mismatch_rows"])
        if blockers:
            raise ValueError("Input failed label/ID validation. Run audit_dataset.py and inspect its report.")
        records = read_records(args.input)
        prepared, excluded, statistics = prepare(records)
        if hashlib.sha256(args.input.read_bytes()).hexdigest() != source_audit["input_sha256"]:
            raise ValueError("Input changed during preparation; close Excel and rerun.")
        args.output_dir.mkdir(parents=True, exist_ok=True)
        dataset_hash = write_atomic(args.output_dir / "news_dataset.jsonl", json_lines(prepared))
        labels_hash = write_atomic(args.output_dir / "labels.json", json.dumps(list(LABEL_NAMES), indent=2) + "\n")
        excluded_hash = write_atomic(args.output_dir / "excluded_records.jsonl", json_lines(excluded))
        manifest = {
            "schema_version": 1,
            "preprocessing_version": PREPROCESSING_VERSION,
            "label_source": "claude_gpt_consensus_silver",
            "target_encoding": {"positive": 1, "provisional_negative": 0, "unknown": -1},
            "unknown_policy": "Exclude unknown article/label pairs from training and evaluation; never convert -1 to 0.",
            "feature_column": "text",
            "label_order_file": "labels.json",
            "input_file": args.input.name,
            "input_sha256": source_audit["input_sha256"],
            "dataset_sha256": dataset_hash,
            "labels_sha256": labels_hash,
            "excluded_records_sha256": excluded_hash,
            "statistics": statistics,
            "limitations": [
                "Consensus targets are not human-validated ground truth.",
                "Only body-based articles are included; title-only articles need a separate policy.",
                "Deduplication covers normalized exact body matches, not near-duplicates or shared events.",
                "No train/validation/test split has been performed yet.",
            ],
        }
        write_atomic(args.output_dir / "preparation_report.json", json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    except (OSError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print(f"Input news: {statistics['input_rows']}")
    print(f"Prepared news: {statistics['prepared_rows']}")
    print(f"Excluded missing body: {statistics['excluded_by_reason'].get('missing_body', 0)}")
    print(f"Removed consistent duplicate extra rows: {statistics['excluded_by_reason'].get('consistent_duplicate', 0)}")
    print(f"Excluded conflicting duplicate rows: {statistics['excluded_by_reason'].get('duplicate_annotation_conflict', 0)}")
    print(f"Rows with no consensus positive: {statistics['rows_with_no_consensus_positive']}")
    print(f"Rows with at least one unknown label: {statistics['rows_with_any_unknown_label']}")
    print(f"Labels: {len(LABEL_NAMES)}")
    print("Label encoding: 1=consensus positive, 0=provisional negative, -1=unknown")
    print(f"Files saved to: {args.output_dir}")
    print("Preparation completed. Next step: freeze the data split.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
