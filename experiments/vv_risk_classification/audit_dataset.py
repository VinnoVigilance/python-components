"""Audit the VV consensus-label sample without changing the source workbook.

Run from the repository root:
    python experiments/vv_risk_classification/audit_dataset.py

Reports model agreement, not accuracy against human-reviewed ground truth.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook


# Names observed in the supplied 23-label sample. Definitions are not inferred.
LABEL_NAMES = (
    "PEP", "RCA", "State-Owned", "Jurisdiction Risk", "Sanctions",
    "Terrorism / Terrorist Financing", "Proliferation / Proliferation Financing",
    "Human Rights / Atrocities", "Regulatory", "Law Enforcement",
    "Judicial / Court", "Crime-Corruption", "Crime-Fraud",
    "Crime-Forgery & Counterfeiting", "Crime-Money Laundering", "Crime-Tax Crime",
    "Crime-Securities & Market Crime", "Crime-Cybercrime",
    "Crime-Trafficking & Smuggling", "Crime-Organized Crime", "Crime-Violent Crime",
    "Crime-Theft & Property Crime", "Crime-Environmental & Wildlife Crime",
)
REQUIRED_COLUMNS = (
    "News ID", "News title", "News body", "Claude tags", "GPT tags",
    "Shared tags (intersection)",
)


def text_value(value: object) -> str:
    return "" if value is None else str(value).strip()


def parse_tags(value: object) -> set[str]:
    if value is None:
        return set()
    if not isinstance(value, str):
        raise ValueError(f"Tag cell must contain text or be empty, got {type(value).__name__}")
    if value.strip().casefold() in {"", "[]", "none", "nan", "no labels"}:
        return set()
    return {part.strip() for part in value.split(",") if part.strip()}


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def audit(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [text_value(value) for value in next(rows, ())]
        missing = sorted(set(REQUIRED_COLUMNS) - set(header))
        repeated = [name for name, count in Counter(header).items() if name and count > 1]
        if missing or repeated:
            raise ValueError(f"Invalid headers. Missing: {missing}; repeated: {repeated}")
        ids = Counter()
        bodies = defaultdict(list)
        annotation_sets = defaultdict(set)
        counts = {name: Counter() for name in ("claude", "gpt", "shared", "disagreement")}
        total = agreed = shared_empty = reviewer_populated = 0
        body_missing, title_missing, id_missing, intersection_mismatch = [], [], [], []
        unknown = Counter()
        unknown_rows = []
        suspicious_lengths = []

        for excel_row, values in enumerate(rows, start=2):
            if not any(value is not None for value in values):
                continue
            total += 1
            record = dict(zip(header, values))
            identifier = text_value(record.get("News ID"))
            row_ref = {"excel_row": excel_row, "news_id": identifier}
            if identifier:
                ids[identifier] += 1
            else:
                id_missing.append(row_ref)
            title = text_value(record.get("News title"))
            body = text_value(record.get("News body"))
            if not title or title.casefold() == "nan":
                title_missing.append(row_ref)
            if not body:
                body_missing.append(row_ref)
            elif len(body) in {32000, 32767}:
                suspicious_lengths.append({**row_ref, "characters": len(body)})

            try:
                claude = parse_tags(record.get("Claude tags"))
                gpt = parse_tags(record.get("GPT tags"))
                shared = parse_tags(record.get("Shared tags (intersection)"))
            except ValueError as error:
                raise ValueError(f"Excel row {excel_row}: {error}") from error
            expected_shared = claude & gpt
            if shared != expected_shared:
                intersection_mismatch.append(row_ref)
            unexpected = (claude | gpt | shared) - set(LABEL_NAMES)
            unknown.update(unexpected)
            if unexpected:
                unknown_rows.append({**row_ref, "labels": sorted(unexpected)})
            for key, labels in (("claude", claude), ("gpt", gpt),
                                ("shared", shared), ("disagreement", claude ^ gpt)):
                counts[key].update(labels)
            agreed += claude == gpt
            shared_empty += not shared
            reviewer_populated += bool(text_value(record.get("reivewer", record.get("reviewer"))))

            if body:
                normalized = re.sub(r"\s+", " ", body).strip().casefold()
                fingerprint = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
                bodies[fingerprint].append(row_ref)
                annotation_sets[fingerprint].add((tuple(sorted(claude)), tuple(sorted(gpt))))

        duplicate_groups = [
            {"body_hash": key, "rows": references,
             "annotation_conflict": len(annotation_sets[key]) > 1}
            for key, references in bodies.items() if len(references) > 1
        ]
        observed = set(counts["claude"]) | set(counts["gpt"]) | set(counts["shared"])
        return {
            "report_schema_version": 1,
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "input_file": path.name,
            "input_sha256": file_hash(path),
            "sheet": sheet.title,
            "headers": header,
            "total_rows": total,
            "expected_label_count": len(LABEL_NAMES),
            "observed_label_count": len(observed),
            "full_agreement_rows": agreed,
            "disagreement_rows": total - agreed,
            "full_agreement_percent": round(100 * agreed / total, 2) if total else None,
            "rows_without_shared_labels": shared_empty,
            "reviewer_name_populated_rows": reviewer_populated,
            "human_final_labels_column_present": any(
                name.casefold() in {"final_labels", "gold_labels"} for name in header
            ),
            "missing_body_rows": body_missing,
            "missing_title_rows": title_missing,
            "missing_id_rows": id_missing,
            "duplicate_news_ids": {key: count for key, count in ids.items() if count > 1},
            "duplicate_nonempty_body_extra_rows": sum(len(group["rows"]) - 1 for group in duplicate_groups),
            "duplicate_body_groups": duplicate_groups,
            "intersection_mismatch_rows": intersection_mismatch,
            "unknown_labels": dict(unknown),
            "unknown_label_rows": unknown_rows,
            "possible_export_length_limit_rows": suspicious_lengths,
            "label_counts": {
                label: {key: counts[key][label] for key in counts} for label in LABEL_NAMES
            },
            "interpretation": "Agreement is not human-validated accuracy. Empty or unreviewed labels are not verified negatives.",
        }
    finally:
        workbook.close()


def main() -> int:
    base = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=base / "data" / "raw" / "News.xlsx")
    parser.add_argument("--output", type=Path, default=base / "outputs" / "dataset_audit.json")
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        parser.error("Input and output must be different files.")
    if not args.input.is_file():
        parser.error(f"Input file not found: {args.input}")
    try:
        report = audit(args.input)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    except (OSError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1

    print(f"Total news: {report['total_rows']}")
    print(f"Observed labels: {report['observed_label_count']}")
    print(f"Full agreement: {report['full_agreement_rows']} ({report['full_agreement_percent']}%)")
    print(f"Disagreement: {report['disagreement_rows']}")
    print(f"Rows without shared labels: {report['rows_without_shared_labels']}")
    print(f"Missing body: {len(report['missing_body_rows'])}")
    print(f"Missing title: {len(report['missing_title_rows'])}")
    print(f"Duplicate nonempty body extra rows: {report['duplicate_nonempty_body_extra_rows']}")
    print(f"Duplicate groups with annotation conflict: {sum(group['annotation_conflict'] for group in report['duplicate_body_groups'])}")
    print(f"Intersection mismatches: {len(report['intersection_mismatch_rows'])}")
    print(f"Unknown labels: {len(report['unknown_labels'])}")
    print(f"Report saved to: {args.output}")
    if not report["total_rows"] or report["missing_id_rows"] or report["duplicate_news_ids"] or report["unknown_labels"] or report["intersection_mismatch_rows"]:
        print("ACTION REQUIRED: inspect dataset_audit.json before preparing training data.")
        return 1
    print("Audit completed. This is an agreement report, not an accuracy evaluation.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())