"""Build the final 900-record gold labels: identical -> kept, verdicts agree -> verdict, else 3-way majority (GPT verdict, Claude verdict, Sonnet)."""
import json
import pandas as pd
from openpyxl import load_workbook

D = "data/reports/watchlist_labeling"
XLSX = f"{D}/watchlist_labels_900_claude_vs_gpt.xlsx"
OUT = f"{D}/watchlist_labels_900_final.jsonl"
LABELS = ["Sanctions", "PEP", "RCA", "Crime-Corruption", "Crime-Fraud", "Crime-Money Laundering",
          "Crime-Trafficking & Smuggling", "Crime-Cybercrime", "Crime-Violent Crime", "Crime-Theft & Property Crime",
          "Crime-Forgery & Counterfeiting", "Crime-Tax Crime", "Crime-Securities & Market Crime", "Crime-Organized Crime",
          "Crime-Environmental & Wildlife Crime", "Terrorism / Terrorist Financing", "Proliferation / Proliferation Financing",
          "Human Rights / Atrocities", "Regulatory", "Law Enforcement", "Judicial / Court", "Jurisdiction Risk", "State-Owned"]


def split(v):
    return set() if pd.isna(v) or str(v).strip() == "" else {s.strip() for s in str(v).split("|") if s.strip()}


def join(s):
    return " | ".join(l for l in LABELS if l in s)


def evidence_map(cell):
    out = {}
    for line in str(cell or "").split("\n"):
        lab = next((l for l in LABELS if line.startswith(l + ": ")), None)
        if lab: out[lab] = line[len(lab) + 2:]
    return out


def decide(r, son):
    cp, cu = split(r.claude_labels), split(r.claude_unsure)
    if r.difference_type == "none":
        return cp, cu, "identical (Claude = GPT)", ""
    gl, gu, ml, mu = split(r.gpt_verdict_labels), split(r.gpt_verdict_unsure), split(r.claude_verdict_labels), split(r.claude_verdict_unsure)
    if gl == ml:
        return gl, (gu if gu == mu else gu | mu) - gl, "both verdicts agree", ""
    s = {d["label"]: d for d in son[r.record_key]["decisions"]}
    pos, uns, notes = gl & ml, (gu | mu) - (gl | ml), []
    for lab in gl ^ ml:
        votes = (lab in gl) + (lab in ml) + (s[lab]["decision"] == "positive")
        (pos if votes >= 2 else uns).add(lab)
        notes.append(f"{lab}: {'added' if votes >= 2 else 'unsure'} ({votes}/3 positive) - {s[lab]['reason']}")
    return pos, uns - pos, "3-way majority (GPT verdict, Claude verdict, Sonnet)", " || ".join(notes)


if __name__ == "__main__":
    df = pd.read_excel(XLSX, sheet_name="Records")
    df = df[[c for c in df.columns if not c.startswith("final_") and not c.endswith("_correct_vs_final")]]
    son = {r["record_key"]: r for r in map(json.loads, open(f"{D}/adjudicate_sonnet_40.jsonl", encoding="utf-8"))}
    manual = {("SECO-SANCTIONS:70496", "PEP"): "Yury Chikhanchin is the Director of Rosfinmonitoring"}
    rows, recs, missing = [], [], 0
    for r in df.itertuples():
        pos, uns, by, note = decide(r, son)
        ev = {**{d["label"]: d["evidence"] for d in son.get(r.record_key, {}).get("decisions", []) if d["evidence"]},
              **evidence_map(r.gpt_evidence), **evidence_map(r.claude_evidence)}
        fev = {l: ev.get(l) or manual.get((r.record_key, l), "") for l in LABELS if l in pos}
        missing += sum(1 for v in fev.values() if not v)
        row = {"final_labels": join(pos), "final_unsure": join(uns), "final_label_count": len(pos), "final_decided_by": by,
               "final_note": note, "final_evidence": "\n".join(f"{l}: {v}" for l, v in fev.items()),
               "claude_correct_vs_final": "Y" if split(r.claude_labels) == pos else "N",
               "gpt_correct_vs_final": "Y" if split(r.gpt_labels) == pos else "N"}
        row.update({f"final_{l}": 1 if l in pos else (-1 if l in uns else 0) for l in LABELS})
        rows.append(row)
        recs.append({"record_key": r.record_key, "source": r.source, "entity_type": r.entity_type, "text": r.input_text,
                     "labels": [l for l in LABELS if l in pos], "unsure": [l for l in LABELS if l in uns],
                     "evidence": fev, "decided_by": by, "taxonomy_version": "VV_Risk_Labels_Final_reviewed"})
    fin = pd.concat([df, pd.DataFrame(rows)], axis=1)
    with open(OUT, "w", encoding="utf-8") as o:
        for x in recs: o.write(json.dumps(x, ensure_ascii=False) + "\n")
    sheets = {ws: pd.read_excel(XLSX, sheet_name=ws) for ws in load_workbook(XLSX, read_only=True).sheetnames if ws != "Records"}
    dist = pd.DataFrame([{"label": l, "final_positive": int((fin[f"final_{l}"] == 1).sum()), "final_unsure": int((fin[f"final_{l}"] == -1).sum())} for l in LABELS])
    with pd.ExcelWriter(XLSX, engine="openpyxl") as w:
        for name, d in sheets.items():
            if name in ("Final Summary", "Final Label Counts"): continue
            d.to_excel(w, sheet_name=name, index=False)
            if name == "Report": fin.to_excel(w, sheet_name="Records", index=False)
        pd.DataFrame([("Records", len(fin), ""), *[(f"Decided by: {k}", v, "") for k, v in fin.final_decided_by.value_counts().items()],
                      ("Records with no final label", int((fin.final_label_count == 0).sum()), ""),
                      ("Total final positive labels", int(fin.final_label_count.sum()), ""),
                      ("Claude original labels = final", int((fin.claude_correct_vs_final == "Y").sum()), f"{(fin.claude_correct_vs_final == 'Y').mean():.1%}"),
                      ("GPT original labels = final", int((fin.gpt_correct_vs_final == "Y").sum()), f"{(fin.gpt_correct_vs_final == 'Y').mean():.1%}"),
                      ("Final labels without an evidence quote", missing, "")], columns=["metric", "count", "pct"]).to_excel(w, sheet_name="Final Summary", index=False)
        dist.to_excel(w, sheet_name="Final Label Counts", index=False)
        for ws in w.book.worksheets:
            ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(50, max(10, max(len(str(c.value or "").split("\n")[0]) for c in col[:200]) + 2))
    print(pd.read_excel(XLSX, sheet_name="Final Summary").to_string(index=False))
    print(dist.to_string(index=False))
