"""Generate the watchlist lookup workbook (list / measure / programme -> risk labels) from lookupRules."""
import csv
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lookupRules import list_rule, measure_rule, program_rule

REPO = Path(__file__).resolve().parents[3]
RESEARCH = Path(__file__).resolve().parents[1]
FINAL = REPO / "data" / "final"
OUT = RESEARCH / "lookups" / "watchlistRiskLookup.xlsx"
SERVER = RESEARCH / "data" / "serverInventory.csv"

NOT_WATCHLIST = {"DNFBP", "ADB_CASE_SUMMARIES", "DTI_FAIR_TRADE_PRESS_RELEASES", "DTI_PH_FAIR_TRADE_PRESS_RELEASES",
                 "NBI_PRESS_RELEASES", "PCIJ_CORRUPTION_WATCH", "PCIJ_INVESTIGATIVE_REPORTS", "PTV_NEWS",
                 "SEC_PH_ADVISORIES", "UK_GOV_NEWS_COMMUNICATIONS"}


def licence_status(note):
    """Pull the status text out of a DMW 'License Status: X' note."""
    m = re.search(r"License Status:\s*(.+)", note or "")
    return m.group(1).strip() if m else ""


def collect():
    """Distinct lists, measures (with DMW licence status) and programme values from local final files and the server export."""
    lists, measures, programs = set(), set(), set()
    with open(SERVER, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["kind"] == "list":
                lists.add(row["list_name"])
            elif row["kind"] == "measure":
                measures.add((row["list_name"], row["type"], row["value"]))
            elif row["kind"] == "program":
                programs.add((row["list_name"], row["type"], row["value"]))
    for f in sorted(FINAL.glob("*_final.jsonl")):
        lst = f.name[: -len("_final.jsonl")]
        if lst in NOT_WATCHLIST:
            continue
        lists.add(lst)
        for line in open(f, encoding="utf-8"):
            r = json.loads(line)
            for m in r.get("Measures") or []:
                status = licence_status(m.get("Note")) if lst == "DMW-RECRUITMENT-AGENCIES" else ""
                measures.add((lst, m.get("MeasureType") or "", status))
            for p in r.get("Programs") or []:
                programs.add((lst, p.get("ProgramType") or "", p.get("Program") or ""))
    return sorted(lists), sorted(measures), sorted(programs)


def write_sheet(wb, title, header, rows, widths):
    """Write one formatted sheet with a header row, filter and frozen pane."""
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="DCE6F0")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    if OUT.exists():
        raise SystemExit(f"{OUT} already exists - it is user-owned now, not regenerating.")
    lists, measures, programs = collect()
    names = json.load(open(RESEARCH / "data" / "label_names.json", encoding="utf-8"))
    tail = ["labels", "basis", "reason", "review_note"]

    list_rows = [[l, " | ".join(r[0]), r[1], r[2], ""] for l in lists for r in [list_rule(l)]]
    measure_rows = []
    for lst, mtype, status in measures:
        labels, basis, reason = measure_rule(lst, mtype, f"License Status: {status}" if status else "")
        measure_rows.append([lst, mtype or "(empty)", status, " | ".join(labels), basis, reason, ""])
    program_rows = []
    for lst, ptype, program in programs:
        rule = program_rule(lst, ptype, program)
        if rule:
            program_rows.append([lst, ptype, program, " | ".join(rule[0]), rule[1], rule[2], ""])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "Lists", ["list_name"] + tail, list_rows, [36, 40, 20, 70, 30])
    write_sheet(wb, "Measures", ["list_name", "measure_type", "licence_status"] + tail, measure_rows, [30, 40, 26, 30, 20, 60, 30])
    write_sheet(wb, "Programs", ["list_name", "program_type", "program"] + tail, program_rows, [30, 16, 55, 40, 20, 60, 30])
    write_sheet(wb, "Labels", ["label"], [[n] for n in names], [45])
    write_sheet(wb, "Basis", ["basis", "meaning"], [
        ["guide", "Already decided in the labelling guide (code table, list rules)."],
        ["data review 2026-10-11", "Decided from the record texts behind the value; reason column says why."],
        ["text decides", "No label from the value alone; the text model decides."],
    ], [24, 90])
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"lists={len(list_rows)} measures={len(measure_rows)} programs={len(program_rows)} -> {OUT}")


if __name__ == "__main__":
    main()
