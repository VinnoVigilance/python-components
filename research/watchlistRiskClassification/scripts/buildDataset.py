"""Build the 900-record watchlist training dataset with one final label set that follows the lookup rules."""
import csv
import json
from pathlib import Path

import openpyxl

from lookupRules import record_rule_labels
from watchlistText import render

REPO = Path(__file__).resolve().parents[3]
FINAL = REPO / "data" / "final"
DATA = Path(__file__).resolve().parents[1] / "data"
LABELLING = DATA / "labelling"
TAXONOMY = DATA / "taxonomy" / "VV_Risk_Labels_Final_reviewed.xlsx"


def split_labels(cell):
    """Turn a 'A | B' cell into a label set."""
    return {x.strip() for x in str(cell or "").split("|") if x.strip()}


def load_label_names():
    """Read the 23 label names from the taxonomy workbook."""
    ws = openpyxl.load_workbook(TAXONOMY, read_only=True)["Risk Labels"]
    return [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def read_csv(path):
    """Read a CSV into a list of dicts (empty list when the file is missing)."""
    if not path.exists():
        return []
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_base_labels():
    """Labels both labellers agreed on, or both verdicts agreed on, or the recorded decision for disputed records."""
    ws = openpyxl.load_workbook(LABELLING / "watchlist_labels_900_claude_vs_gpt.xlsx", read_only=True)["Records"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    decided = {d["record_key"]: split_labels(d["final_labels"]) for d in read_csv(DATA / "disputed_decisions.csv")}
    base = {}
    for r in rows:
        c = dict(zip(header, r))
        if c["verdicts_same_labels"] is None:
            base[c["record_key"]] = split_labels(c["labels_both"])
        elif c["verdicts_same_labels"] == "Y":
            base[c["record_key"]] = split_labels(c["claude_verdict_labels"])
        else:
            base[c["record_key"]] = decided[c["record_key"]]
    return base


def find_canonical(inputs):
    """Map record_key -> canonical record, preferring the one whose rendered text equals the sample text."""
    wanted = {}
    for r in inputs:
        wanted.setdefault(r["source"], {})[r["record_key"].split(":", 1)[1]] = r
    exact, same_id = {}, {}
    for lst, by_id in wanted.items():
        for line in open(FINAL / f"{lst}_final.jsonl", encoding="utf-8"):
            rec = json.loads(line)
            sample = by_id.get(str(rec.get("EntityId")))
            if sample:
                same_id.setdefault(sample["record_key"], rec)
                if render(rec, lst) == sample["text"]:
                    exact[sample["record_key"]] = rec
    return {**same_id, **exact}


def main():
    names = load_label_names()
    base = load_base_labels()
    with open(LABELLING / "watchlist_label_sample_900.csv", encoding="utf-8-sig", newline="") as f:
        inputs = list(csv.DictReader(f))
    canonical = find_canonical(inputs)
    missing = [r["record_key"] for r in inputs if r["record_key"] not in canonical]
    assert not missing, f"no canonical record for {missing[:5]}"

    corrections = read_csv(DATA / "label_corrections.csv")
    changes, records = [], []
    for r in inputs:
        key = r["record_key"]
        labels = set(base[key])
        for label, rule in record_rule_labels(r["source"], canonical[key]).items():
            if label not in labels:
                labels.add(label)
                changes.append([key, label, "add", f"lookup rule: {rule}"])
        for c in (c for c in corrections if c["record_key"] == key):
            (labels.add if c["action"] == "add" else labels.discard)(c["label"])
            changes.append([key, c["label"], c["action"], c["reason"]])
        assert labels <= set(names), labels - set(names)
        records.append({"record_key": key, "source": r["source"], "entity_type": r["entity_type"],
                        "part": r["part"], "text": r["text"], "labels": [n for n in names if n in labels]})

    with open(DATA / "dataset_900.jsonl", "w", encoding="utf-8") as f:
        for rec in records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    with open(DATA / "labels_900_binary.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["record_key", "source", "part"] + names)
        for rec in records:
            w.writerow([rec["record_key"], rec["source"], rec["part"]] + [int(n in rec["labels"]) for n in names])
    with open(DATA / "label_changes.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["record_key", "label", "action", "reason"])
        w.writerows(changes)
    with open(DATA / "label_names.json", "w", encoding="utf-8") as f:
        json.dump(names, f, ensure_ascii=False, indent=1)
    print(f"records={len(records)} label changes={len(changes)}")


if __name__ == "__main__":
    main()
