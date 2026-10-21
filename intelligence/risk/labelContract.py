"""Read the business workbook and validate a model's ordered label columns."""
from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[2]


class RiskModelError(RuntimeError):
    """The configured model is missing, incompatible, or invalid."""


@dataclass(frozen=True)
class ModelLabelContract:
    model_id: str
    run_id: str
    taxonomy_version: str
    label_ids: tuple[str, ...]
    artifact_labels: tuple[str, ...]
    display_labels: tuple[str, ...]


def _text(value, field: str) -> str:
    if not isinstance(value, str) or not value.strip() or value != value.strip():
        raise RiskModelError(f"{field} must be a non-empty, trimmed string.")
    return value


def _read_json(path: Path) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as error:
        raise RiskModelError(f"Cannot read {path}: {error}") from error
    if not isinstance(value, dict):
        raise RiskModelError(f"{path} must contain a JSON object.")
    if type(value.get("schema_version")) is not int or value["schema_version"] != 1:
        raise RiskModelError(f"Unsupported schema_version in {path}.")
    return value


def _parse_taxonomy_rows(risk_rows: list[tuple], annotation_rows: list[tuple]) -> dict:
    """Join label definitions to the workbook's existing stable IDs by name."""
    def header(rows, required, sheet):
        if not rows:
            raise RiskModelError(f"Taxonomy sheet {sheet!r} is empty.")
        names = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        if any(names.count(name) != 1 for name in required):
            raise RiskModelError(f"{sheet!r} requires unique columns: {', '.join(required)}.")
        return {name: names.index(name) for name in required}

    risk_columns = header(risk_rows, ("Risk Label", "English Description"), "Risk Labels")
    annotation_columns = header(annotation_rows, ("Stable Label ID", "Risk Label"), "Annotation Rules")
    definitions = {}
    for index, row in enumerate(risk_rows[1:], start=2):
        name = row[risk_columns["Risk Label"]]
        if name is None or name == "":
            continue
        name = _text(name, f"Risk Labels!row {index} label")
        if name in definitions:
            raise RiskModelError(f"Duplicate Risk Label {name!r} at row {index}.")
        definitions[name] = _text(row[risk_columns["English Description"]],
                                  f"Risk Labels!row {index} description")
    if not definitions:
        raise RiskModelError("Risk Labels contains no label definitions.")

    by_id, seen_names = {}, set()
    for index, row in enumerate(annotation_rows[1:], start=2):
        name = row[annotation_columns["Risk Label"]]
        if name is None or name == "":
            continue
        name = _text(name, f"Annotation Rules!row {index} label")
        label_id = _text(row[annotation_columns["Stable Label ID"]],
                         f"Annotation Rules!row {index} Stable Label ID")
        if name not in definitions:
            raise RiskModelError(f"Annotation Rules label {name!r} has no Risk Labels definition.")
        if label_id in by_id or name in seen_names:
            raise RiskModelError(f"Duplicate stable ID or label at Annotation Rules row {index}.")
        by_id[label_id] = {"display_name": name, "definition": definitions[name]}
        seen_names.add(name)
    missing = set(definitions) - seen_names
    if missing:
        raise RiskModelError(f"Missing stable IDs for Risk Labels: {', '.join(sorted(missing))}.")
    return by_id


def _read_workbook_taxonomy(data: bytes) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        required = {"Risk Labels", "Annotation Rules"}
        if not required.issubset(workbook.sheetnames):
            raise RiskModelError("Taxonomy workbook needs Risk Labels and Annotation Rules sheets.")
        return _parse_taxonomy_rows(
            list(workbook["Risk Labels"].iter_rows(values_only=True)),
            list(workbook["Annotation Rules"].iter_rows(values_only=True)),
        )
    finally:
        workbook.close()


def _taxonomy_fingerprint(labels: dict) -> str:
    content = json.dumps(labels, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return "sha256:" + hashlib.sha256(content.encode("utf-8")).hexdigest()


def load_taxonomy(path=None):
    """Read business labels and fingerprint their semantic content."""
    path = Path(path) if path else PROJECT_ROOT / "data/rules/VV_RiskCategory_Labels.xlsx"
    data = path.read_bytes()
    labels = _read_workbook_taxonomy(data)
    return _taxonomy_fingerprint(labels), labels


def load_label_contract(directory: Path) -> ModelLabelContract:
    manifest = _read_json(directory / "manifest.json")
    model_id = _text(manifest.get("model_id"), "model_id")
    run_id = _text(manifest.get("run_id"), "run_id")
    filename = _text(manifest.get("taxonomy_file"), "taxonomy_file")
    relative = Path(filename)
    taxonomy_path = (PROJECT_ROOT / relative).resolve()
    if relative.is_absolute() or "\\" in filename or not taxonomy_path.is_relative_to(PROJECT_ROOT.resolve()):
        raise RiskModelError("taxonomy_file must be a project-relative path inside the project.")
    try:
        data = taxonomy_path.read_bytes()
        by_id = _read_workbook_taxonomy(data)
    except RiskModelError:
        raise
    except Exception as error:
        raise RiskModelError(f"Cannot read taxonomy workbook {taxonomy_path}: {error}") from error
    taxonomy_version = _taxonomy_fingerprint(by_id)

    columns = manifest.get("labels")
    if not isinstance(columns, list) or not columns:
        raise RiskModelError("Manifest labels must be a non-empty ordered list.")
    ids, artifact_labels, display_labels = [], [], []
    for column in columns:
        if not isinstance(column, dict):
            raise RiskModelError("Each model label column must be an object.")
        label_id = _text(column.get("id"), "model label id")
        artifact_name = _text(column.get("artifact_name"), "artifact_name")
        if label_id not in by_id:
            raise RiskModelError(f"Unknown taxonomy label id: {label_id}.")
        if artifact_name != by_id[label_id]["display_name"]:
            raise RiskModelError(f"Model label {artifact_name!r} does not match taxonomy id {label_id}.")
        if label_id in ids or artifact_name in artifact_labels:
            raise RiskModelError("Model label ids and artifact names must be unique.")
        ids.append(label_id)
        artifact_labels.append(artifact_name)
        display_labels.append(by_id[label_id]["display_name"])
    if len(set(display_labels)) != len(display_labels):
        raise RiskModelError("Supported labels must have distinct display names.")
    return ModelLabelContract(model_id, run_id, taxonomy_version,
                              tuple(ids), tuple(artifact_labels), tuple(display_labels))
