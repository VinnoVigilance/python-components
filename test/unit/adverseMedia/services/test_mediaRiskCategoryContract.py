"""Model-label contract tests; no DB, network, or original pickle loading.

Run with the standard library, even when pytest is not installed:
python -m unittest discover -s test/unit/adverseMedia/services -p test_mediaRiskCategoryContract.py -v
"""
from __future__ import annotations

import copy
import shutil
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
from sklearn.feature_extraction.text import CountVectorizer
from sklearn.svm import LinearSVC

from intelligence.risk.mediaRiskCategory import nbsvm
from intelligence.risk import labelContract as contracts
from intelligence.risk.labelContract import load_label_contract, RiskModelError as MediaRiskModelError
from intelligence.risk.mediaRiskCategory.predictor import MediaRiskPredictor


ROOT = Path(__file__).resolve().parents[4]
BUNDLE = ROOT / "intelligence/risk/mediaRiskCategory/models/nbsvm-v1"


class MediaRiskContractTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.project = Path(self.temporary.name)
        self.directory = self.project / 'models/nbsvm-v1'
        self.directory.mkdir(parents=True)
        for name in ('run.json', 'manifest.json'):
            (self.directory / name).write_bytes((BUNDLE / name).read_bytes())
        self.workbook_path = self.project / 'data/rules/VV_RiskCategory_Labels.xlsx'
        self.workbook_path.parent.mkdir(parents=True)
        shutil.copyfile(ROOT / 'data/rules/VV_RiskCategory_Labels.xlsx', self.workbook_path)
        patcher = patch.object(contracts, 'PROJECT_ROOT', self.project)
        patcher.start()
        self.addCleanup(patcher.stop)
        (self.directory / 'model.joblib').write_bytes(b'synthetic test placeholder')

    def taxonomy_rows(self):
        from openpyxl import load_workbook
        book = load_workbook(self.workbook_path, read_only=True, data_only=False)
        try:
            return ([list(row) for row in book['Risk Labels'].iter_rows(values_only=True)],
                    [list(row) for row in book['Annotation Rules'].iter_rows(values_only=True)])
        finally:
            book.close()

    def read(self, name):
        return json.loads((self.directory / name).read_text(encoding="utf-8"))

    def write(self, name, value):
        (self.directory / name).write_text(json.dumps(value), encoding="utf-8")

    def alter(self, name, change):
        value = self.read(name)
        change(value)
        self.write(name, value)

    def source(self):
        metadata = self.read("run.json")
        vectorizer = CountVectorizer(binary=True, ngram_range=(1, 2))
        features = vectorizer.fit_transform(["court fraud arrest", "garden flower day"])
        svm = LinearSVC(random_state=42).fit(features, [1, 0])
        classifier = nbsvm._NBSVMState()
        count = len(metadata["labels"])
        classifier.estimators_ = [copy.deepcopy(svm) for _ in range(count)]
        classifier.log_count_ratios_ = np.ones((count, features.shape[1]))
        wrapper = nbsvm._ThresholdState()
        wrapper.estimator_ = SimpleNamespace(named_steps={"counts": vectorizer, "classifier": classifier})
        wrapper.decision_thresholds_ = np.array(metadata["thresholds"])
        return {"model": wrapper, "labels": metadata["labels"],
                "thresholds": metadata["thresholds"], "score_method": "decision_function",
                "dataset_id": metadata["dataset_id"], "protocol_id": metadata["protocol_id"]}

    def load(self, source=None):
        if source is None:
            source = self.source()
        with patch.object(nbsvm, "_check_environment"), patch.object(nbsvm, "_load_original", return_value=source):
            return nbsvm.load_model(self.directory)

    def service(self):
        model, metadata = self.load()
        with patch("intelligence.risk.mediaRiskCategory.predictor.load_model", return_value=(model, metadata)):
            return MediaRiskPredictor(self.directory)

    def test_manifest_preserves_original_label_order_and_names(self):
        contract = load_label_contract(self.directory)
        self.assertEqual(list(contract.artifact_labels), self.read("run.json")["labels"])
        self.assertEqual(contract.display_labels, contract.artifact_labels)
        self.assertEqual(len(contract.label_ids), 11)

    def test_taxonomy_entry_order_does_not_change_model_columns(self):
        original = load_label_contract(self.directory)
        risk, annotations = self.taxonomy_rows()
        risk[1:] = reversed(risk[1:])
        annotations[1:] = reversed(annotations[1:])
        altered = contracts._parse_taxonomy_rows(risk, annotations)
        with patch.object(contracts, '_read_workbook_taxonomy', return_value=altered):
            self.assertEqual(load_label_contract(self.directory), original)
            self.load()

    def test_reordered_manifest_is_rejected_before_model_read(self):
        self.alter("manifest.json", lambda x: x["labels"].reverse())
        with patch.object(nbsvm, "_load_original") as reader:
            with self.assertRaises(MediaRiskModelError):
                nbsvm.load_model(self.directory)
            reader.assert_not_called()

    def test_unknown_label_id_rejected(self):
        self.alter("manifest.json", lambda x: x["labels"][0].update(id="unknown"))
        with self.assertRaisesRegex(MediaRiskModelError, "Unknown taxonomy"):
            load_label_contract(self.directory)

    def test_duplicate_manifest_id_rejected(self):
        self.alter("manifest.json", lambda x: x["labels"][1].update(id=x["labels"][0]["id"]))
        with self.assertRaises(MediaRiskModelError):
            load_label_contract(self.directory)

    def test_duplicate_artifact_name_rejected(self):
        self.alter("manifest.json", lambda x: x["labels"][1].update(artifact_name=x["labels"][0]["artifact_name"]))
        with self.assertRaises(MediaRiskModelError):
            load_label_contract(self.directory)

    def test_duplicate_taxonomy_id_rejected(self):
        risk, annotations = self.taxonomy_rows()
        annotations[2][0] = annotations[1][0]
        with self.assertRaisesRegex(MediaRiskModelError, 'Duplicate stable ID'):
            contracts._parse_taxonomy_rows(risk, annotations)

    def test_taxonomy_version_tracks_label_content(self):
        original = load_label_contract(self.directory)
        parsed = contracts._read_workbook_taxonomy(self.workbook_path.read_bytes())
        parsed['SAN_GEN']['definition'] = 'Revised sanctions definition'
        with patch.object(contracts, '_read_workbook_taxonomy', return_value=parsed):
            changed = load_label_contract(self.directory)
        self.assertNotEqual(original.taxonomy_version, changed.taxonomy_version)
        self.assertEqual(original.label_ids, changed.label_ids)

    def test_missing_manifest_is_actionable(self):
        (self.directory / "manifest.json").unlink()
        with self.assertRaisesRegex(MediaRiskModelError, "manifest.json"):
            self.load()

    def test_missing_taxonomy_is_actionable(self):
        self.workbook_path.unlink()
        with self.assertRaisesRegex(MediaRiskModelError, 'VV_RiskCategory_Labels.xlsx'):
            self.load()

    def test_invalid_json_rejected(self):
        (self.directory / "manifest.json").write_text("{broken", encoding="utf-8")
        with self.assertRaises(MediaRiskModelError):
            load_label_contract(self.directory)

    def test_unknown_schema_version_rejected(self):
        self.alter("manifest.json", lambda x: x.update(schema_version=2))
        with self.assertRaisesRegex(MediaRiskModelError, "schema_version"):
            load_label_contract(self.directory)

    def test_taxonomy_outside_project_rejected(self):
        for filename in ('../outside.xlsx', '/tmp/outside.xlsx', '..\\outside.xlsx'):
            with self.subTest(filename=filename):
                self.alter('manifest.json', lambda x: x.update(taxonomy_file=filename))
                with self.assertRaises(MediaRiskModelError):
                    load_label_contract(self.directory)

    def test_run_identity_mismatch_rejected(self):
        self.alter("manifest.json", lambda x: x.update(run_id="other_run"))
        with self.assertRaisesRegex(MediaRiskModelError, "run.json"):
            self.load()

    def test_artifact_label_order_mismatch_rejected(self):
        source = self.source()
        source["labels"] = list(reversed(source["labels"]))
        with self.assertRaisesRegex(MediaRiskModelError, "same model"):
            self.load(source)

    def test_existing_threshold_validation_preserved(self):
        source = self.source()
        source["thresholds"] = [99.0] * len(source["labels"])
        with self.assertRaisesRegex(MediaRiskModelError, "thresholds"):
            self.load(source)

    def test_existing_dimension_validation_preserved(self):
        source = self.source()
        source["model"].estimator_.named_steps["classifier"].log_count_ratios_ = np.ones((1, 1))
        with self.assertRaisesRegex(MediaRiskModelError, "dimensions"):
            self.load(source)

    def test_23_business_labels_do_not_expand_11_model_columns(self):
        parsed = contracts._read_workbook_taxonomy(self.workbook_path.read_bytes())
        self.assertEqual(len(parsed), 23)
        self.assertEqual(len(self.service().supported_labels), 11)
        self.assertEqual(self.service().supported_label_ids[0], 'TER_GEN')
        self.assertNotIn('SAN_GEN', self.service().supported_label_ids)

    def test_workbook_display_name_mismatch_requires_model_manifest_update(self):
        source = self.source()
        risk, annotations = self.taxonomy_rows()
        original_name = source['labels'][0]
        for row in risk[1:]:
            if row[0] == original_name:
                row[0] = 'Revised display name'
        for row in annotations[1:]:
            if row[0] == 'TER_GEN':
                row[1] = 'Revised display name'
        parsed = contracts._parse_taxonomy_rows(risk, annotations)
        with patch.object(contracts, '_read_workbook_taxonomy', return_value=parsed):
            with self.assertRaises(MediaRiskModelError):
                load_label_contract(self.directory)

    def test_adapter_derives_output_width_from_model_manifest(self):
        for filename in ("manifest.json", "run.json"):
            self.alter(filename, lambda x: x.update(labels=x["labels"][:2]))
        self.alter("run.json", lambda x: x.update(thresholds=x["thresholds"][:2]))
        model, _ = self.load()
        self.assertEqual(nbsvm.decision_scores(model, ["court fraud"]).shape, (1, 2))
        self.assertEqual(nbsvm.decision_scores(model, []).shape, (0, 2))

    def test_original_scoring_and_label_decisions_preserved(self):
        model, _ = self.load()
        texts = [nbsvm.prepare_text("Title", "court fraud arrest"), nbsvm.prepare_text(None, "garden flower day")]
        features = model["vectorizer"].transform(texts)
        previous_scores = np.column_stack([svm.decision_function(features.multiply(ratio).tocsr())
                                          for svm, ratio in zip(model["classifiers"], model["ratios"])])
        np.testing.assert_array_equal(nbsvm.decision_scores(model, texts), previous_scores)
        service = self.service()
        results = service.classify_many([("Title", "court fraud arrest"), (None, "garden flower day")])
        names = self.read("run.json")["labels"]
        for values, result in zip(previous_scores, results):
            expected = [label for label, value, threshold in zip(names, values, model["thresholds"]) if value >= threshold]
            self.assertEqual(result["labels"], expected)

    def test_threshold_tie_remains_inclusive(self):
        service = self.service()
        values = np.array([service._model["thresholds"]])
        with patch("intelligence.risk.mediaRiskCategory.predictor.decision_scores", return_value=values):
            result = service.classify("Title", "Body")
        self.assertEqual(tuple(result["labels"]), service.supported_labels)
        self.assertEqual(tuple(result["label_ids"]), service.supported_label_ids)

    def test_empty_body_batch_order_and_provenance(self):
        service = self.service()
        articles = [(None, None), ("Title", "court fraud"), ("Title", " ")]
        batch = service.classify_many(articles)
        self.assertEqual(batch, [service.classify(*article) for article in articles])
        self.assertEqual(service.classify_many([]), [])
        self.assertEqual([r["status"] for r in batch], ["SKIPPED_EMPTY_BODY", "CLASSIFIED", "SKIPPED_EMPTY_BODY"])
        self.assertEqual(batch[0]["label_ids"], [])
        for result in batch:
            self.assertEqual(result["taxonomy_version"], load_label_contract(self.directory).taxonomy_version)
            self.assertEqual(result["model_version"], "media-risk-nbsvm-v1:" + self.read("run.json")["run_id"])

    def test_preparation_type_and_normalization_unchanged(self):
        self.assertEqual(nbsvm.prepare_text("Ｔｅｓｔ", " A\n B "), "Test\n\nA B")
        with self.assertRaises(TypeError):
            nbsvm.prepare_text("Title", 42)

    def test_environment_guard_still_enforced(self):
        with self.assertRaises(MediaRiskModelError):
            nbsvm._check_environment({"python": "0.0.0"})

    def test_single_requirements_pins_match_selected_run_environment(self):
        pins = dict(line.split('#', 1)[0].strip().split('==')
                    for line in (ROOT / 'requirements.txt').read_text().splitlines()
                    if line.strip() and not line.startswith('#') and '==' in line)
        environment = self.read('run.json')['environment']
        for package, field in (('numpy', 'numpy'), ('scipy', 'scipy'),
                               ('scikit-learn', 'scikit_learn'), ('joblib', 'joblib')):
            self.assertEqual(pins[package], environment[field])
        self.assertFalse((ROOT / 'requirements-media-risk.txt').exists())
        self.assertNotIn('requirements-media-risk', (ROOT / 'Dockerfile').read_text())
        self.assertIn('COPY requirements.txt .', (ROOT / 'Dockerfile').read_text())


    def test_missing_workbook_stable_id_rejected(self):
        risk, annotations = self.taxonomy_rows()
        del annotations[1]
        with self.assertRaisesRegex(MediaRiskModelError, 'Missing stable IDs'):
            contracts._parse_taxonomy_rows(risk, annotations)

    def test_duplicate_business_label_rejected(self):
        risk, annotations = self.taxonomy_rows()
        risk.append(risk[1])
        with self.assertRaisesRegex(MediaRiskModelError, 'Duplicate Risk Label'):
            contracts._parse_taxonomy_rows(risk, annotations)

    def test_missing_workbook_headers_rejected(self):
        risk, annotations = self.taxonomy_rows()
        annotations[0][0] = 'Wrong header'
        with self.assertRaisesRegex(MediaRiskModelError, 'requires unique columns'):
            contracts._parse_taxonomy_rows(risk, annotations)

    def test_shared_guidance_rows_are_not_labels(self):
        risk, annotations = self.taxonomy_rows()
        parsed = contracts._parse_taxonomy_rows(risk, annotations)
        self.assertEqual(len(parsed), 23)
        self.assertNotIn('G01 Scope', parsed)


if __name__ == "__main__":
    unittest.main()
